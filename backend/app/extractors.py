import io
from pathlib import Path


def extract_resume(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError('PDF support is not installed on the server')
        reader = PdfReader(io.BytesIO(content))
        text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        if not text.strip():
            raise ValueError('No text was extracted. This appears to be a scanned PDF; OCR is not currently enabled.')
        return text.strip()
    if suffix == '.docx':
        try:
            import docx
        except ImportError:
            raise ValueError('DOCX support is not installed on the server')
        doc = docx.Document(io.BytesIO(content))
        text_parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                text_parts.extend(cell.text for cell in row.cells)
        return '\n'.join(text_parts).strip()
    raise ValueError('Unsupported file type. Please upload a PDF or DOCX file.')
